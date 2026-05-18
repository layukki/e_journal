import io
import sqlite3
import hashlib
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="Электронный журнал", layout="wide", initial_sidebar_state="expanded"
)

conn = sqlite3.connect("journal.db", check_same_thread=False)
c = conn.cursor()

# ======================= СОЗДАНИЕ ТАБЛИЦ =======================
c.execute("""
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE,
    password TEXT,
    full_name TEXT,
    role TEXT,
    group_id INTEGER
)
""")

c.execute("""
CREATE TABLE IF NOT EXISTS parent_students (
    parent_id INTEGER,
    student_id INTEGER,
    PRIMARY KEY (parent_id, student_id)
)
""")

c.execute("""
CREATE TABLE IF NOT EXISTS groups (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT UNIQUE
)
""")
c.execute("""
CREATE TABLE IF NOT EXISTS disciplines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT
)
""")
c.execute("""
CREATE TABLE IF NOT EXISTS group_disciplines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    group_id INTEGER,
    discipline_id INTEGER,
    teacher_id INTEGER,
    UNIQUE(group_id, discipline_id, teacher_id)
)
""")
c.execute("""
CREATE TABLE IF NOT EXISTS lessons (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    group_discipline_id INTEGER,
    date TEXT,
    homework TEXT,
    topic TEXT
)
""")
c.execute("""
CREATE TABLE IF NOT EXISTS grades (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    lesson_id INTEGER,
    student_id INTEGER,
    grade TEXT,
    UNIQUE(lesson_id, student_id)
)
""")
conn.commit()


# ======================= ЭКСПОРТ В EXCEL (ИСПРАВЛЕНО) =======================
def generate_excel(
    disc_name,
    group_name,
    teacher_name,
    dates,
    students,
    lesson_map,
    topic_map,
    hw_map,
    conn_db,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал"

    header_font_white = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="4472C4", fgColor="4472C4")
    meta_fill = PatternFill("solid", start_color="D9E1F2", fgColor="D9E1F2")
    meta_font = Font(name="Arial", bold=True, size=10)
    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    date_cols = [pd.to_datetime(d).strftime("%d.%m") for d in dates]
    total_cols = 2 + len(dates)

    # Заголовок
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(
        row=1, column=1, value=f"Журнал: {disc_name}  •  Группа {group_name}"
    )
    title_cell.font = Font(name="Arial", bold=True, size=13)
    title_cell.alignment = center

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    sub_cell = ws.cell(row=2, column=1, value=f"Преподаватель: {teacher_name}")
    sub_cell.font = Font(name="Arial", italic=True, size=10)
    sub_cell.alignment = left_wrap

    # Шапка
    HEADER_ROW = 4
    for col_idx, val in enumerate(["№", "ФИО студента"], start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=val)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for j, dcol in enumerate(date_cols, start=3):
        cell = ws.cell(row=HEADER_ROW, column=j, value=dcol)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # Студенты + Оценки с правильным типом
    cur = conn_db.cursor()
    for i, (sid, name) in enumerate(students, 1):
        row_idx = HEADER_ROW + i

        # №
        ws.cell(row_idx, 1, i).font = cell_font
        ws.cell(row_idx, 1).alignment = center
        ws.cell(row_idx, 1).border = border

        # ФИО
        ws.cell(row_idx, 2, name).font = cell_font
        ws.cell(row_idx, 2).alignment = left_wrap
        ws.cell(row_idx, 2).border = border

        for j, d in enumerate(dates, start=3):
            g = cur.execute(
                "SELECT grade FROM grades WHERE lesson_id=? AND student_id=?",
                (lesson_map[d], sid),
            ).fetchone()

            grade_val = g[0] if g and g[0] else ""

            cell = ws.cell(row=row_idx, column=j)
            cell.font = cell_font
            cell.alignment = center
            cell.border = border

            if grade_val in ["2", "3", "4", "5"]:
                cell.value = int(grade_val)
                cell.data_type = "n"
            elif grade_val == "Н":
                cell.value = "Н"
                cell.data_type = "s"
            else:
                cell.value = ""

    # Тема занятия
    topic_row_idx = HEADER_ROW + len(students) + 1
    ws.cell(row=topic_row_idx, column=1, value="").border = border
    tc = ws.cell(row=topic_row_idx, column=2, value="Тема занятия")
    tc.font = meta_font
    tc.fill = meta_fill
    tc.alignment = left_wrap
    tc.border = border

    for j, d in enumerate(dates, start=3):
        cell = ws.cell(row=topic_row_idx, column=j, value=topic_map.get(d, ""))
        cell.font = cell_font
        cell.fill = meta_fill
        cell.alignment = left_wrap
        cell.border = border

    # Домашнее задание
    hw_row_idx = topic_row_idx + 1
    ws.cell(row=hw_row_idx, column=1, value="").border = border
    hc = ws.cell(row=hw_row_idx, column=2, value="Домашнее задание")
    hc.font = meta_font
    hc.fill = meta_fill
    hc.alignment = left_wrap
    hc.border = border

    for j, d in enumerate(dates, start=3):
        cell = ws.cell(row=hw_row_idx, column=j, value=hw_map.get(d, ""))
        cell.font = cell_font
        cell.fill = meta_fill
        cell.alignment = left_wrap
        cell.border = border

    # Ширина колонок
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    for j in range(3, 3 + len(dates)):
        ws.column_dimensions[get_column_letter(j)].width = 9

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ======================= АВТОРИЗАЦИЯ И ОСНОВНАЯ ЛОГИКА =======================
if "user" not in st.session_state:
    st.title("Электронный журнал")
    with st.form("login"):
        username = st.text_input("Логин")
        password = st.text_input("Пароль", type="password")
        if st.form_submit_button("Войти"):
            h = hashlib.sha256(password.encode()).hexdigest()
            user_data = c.execute(
                "SELECT id,full_name,role,group_id FROM users WHERE username=? AND password=?",
                (username, h),
            ).fetchone()
            if user_data:
                st.session_state.user = {
                    "id": user_data[0],
                    "full_name": user_data[1],
                    "role": user_data[2],
                    "group_id": user_data[3],
                }
                st.rerun()
            else:
                st.error("Неверный логин или пароль")
else:
    user = st.session_state.user
    st.sidebar.header(user["full_name"])

    role_name = {
        "admin": "Администратор",
        "teacher": "Преподаватель",
        "student": "Студент",
        "parent": "Родитель",
    }.get(user["role"], user["role"])

    st.sidebar.caption(f"Роль: {role_name}")

    if st.sidebar.button("Выйти"):
        del st.session_state.user
        st.rerun()

    menu = ["Журнал"]
    if user["role"] == "admin":
        menu += ["Пользователи", "Группы", "Предметы", "Назначение"]

    page = st.sidebar.radio("Меню", menu)

    # ======================= ЖУРНАЛ =======================
    if page == "Журнал":
        st.header("Журнал")

        if user["role"] == "parent":
            st.info(f"👨‍👩‍👧 **Залогинен как родитель** — {user['full_name']}")

            children = c.execute(
                """
                SELECT u.id, u.full_name, u.group_id 
                FROM parent_students ps
                JOIN users u ON u.id = ps.student_id
                WHERE ps.parent_id = ?
                ORDER BY u.full_name
            """,
                (user["id"],),
            ).fetchall()

            if not children:
                st.warning("У вас пока не привязано ни одного ученика.")
                st.stop()

            if len(children) > 1:
                child_names = [ch[1] for ch in children]
                selected_child_name = st.selectbox("Выберите ребёнка", child_names)
                child = next(ch for ch in children if ch[1] == selected_child_name)
            else:
                child = children[0]

            effective_student_id = child[0]
            effective_group_id = child[2]

        else:
            effective_student_id = user["id"]
            effective_group_id = user.get("group_id")

        if user["role"] in ["student", "parent"]:
            sql, params = "WHERE gd.group_id = ?", (effective_group_id,)
        elif user["role"] == "teacher":
            sql, params = "WHERE gd.teacher_id = ?", (user["id"],)
        else:
            sql, params = "", ()

        assignments = c.execute(
            f"""
            SELECT gd.id, g.name, d.name, u.full_name, gd.group_id
            FROM group_disciplines gd
            JOIN groups g ON g.id = gd.group_id
            JOIN disciplines d ON d.id = gd.discipline_id
            JOIN users u ON u.id = gd.teacher_id
            {sql}
            ORDER BY d.name, g.name
            """,
            params,
        ).fetchall()

        if not assignments:
            st.info("Нет доступных журналов")
        else:
            options = [f"{d} — {g} (преп. {t})" for _, g, d, t, _ in assignments]
            sel = st.selectbox("Выберите журнал", options)
            idx = options.index(sel)
            assignment_id, group_name, disc_name, teacher_name, group_id = assignments[
                idx
            ]

            st.subheader(f"{disc_name} • Группа {group_name}")

            lessons = c.execute(
                "SELECT id, date, homework, topic FROM lessons WHERE group_discipline_id=? ORDER BY date",
                (assignment_id,),
            ).fetchall()

            if not lessons:
                st.info("Пока нет уроков по этому предмету")
            else:
                dates = [l[1] for l in lessons]
                lesson_map = {l[1]: l[0] for l in lessons}
                hw_map = {l[1]: (l[2] or "") for l in lessons}
                topic_map = {l[1]: (l[3] or "") for l in lessons}

                if user["role"] in ["student", "parent"]:
                    st.markdown("### Мои оценки")
                    grades_data = {"Дата": [], "Оценка": []}
                    for d in dates:
                        g = c.execute(
                            "SELECT grade FROM grades WHERE lesson_id=? AND student_id=?",
                            (lesson_map[d], effective_student_id),
                        ).fetchone()
                        grades_data["Дата"].append(
                            pd.to_datetime(d).strftime("%d.%m.%Y")
                        )
                        grades_data["Оценка"].append(g[0] if g and g[0] else "")

                    st.dataframe(
                        pd.DataFrame(grades_data),
                        use_container_width=True,
                        hide_index=True,
                    )

                    st.markdown("### Уроки и задания")
                    lesson_data = [
                        {
                            "Дата": pd.to_datetime(d).strftime("%d.%m.%Y"),
                            "Тема занятия": topic_map.get(d, "—"),
                            "Домашнее задание": hw_map.get(d, "—"),
                        }
                        for d in dates
                    ]
                    st.dataframe(
                        pd.DataFrame(lesson_data),
                        use_container_width=True,
                        hide_index=True,
                    )

                else:
                    students = c.execute(
                        "SELECT id, full_name FROM users WHERE group_id=? AND role='student' ORDER BY full_name",
                        (group_id,),
                    ).fetchall()

                    date_cols = [pd.to_datetime(d).strftime("%d.%m") for d in dates]

                    grade_columns = ["№", "ФИО студента"] + date_cols
                    grade_data = []
                    for i, (sid, name) in enumerate(students, 1):
                        row = {"№": i, "ФИО студента": name}
                        for d in dates:
                            g = c.execute(
                                "SELECT grade FROM grades WHERE lesson_id=? AND student_id=?",
                                (lesson_map[d], sid),
                            ).fetchone()
                            row[pd.to_datetime(d).strftime("%d.%m")] = (
                                str(g[0]).strip() if g and g[0] is not None else ""
                            )
                        grade_data.append(row)

                    grade_df = pd.DataFrame(grade_data, columns=grade_columns)
                    grade_col_config = {
                        col: st.column_config.SelectboxColumn(
                            col,
                            options=["", "2", "3", "4", "5", "Н"],
                            required=False,
                            default="",
                        )
                        for col in date_cols
                    }

                    st.markdown("**Оценки**")
                    edited_grade_df = st.data_editor(
                        grade_df,
                        use_container_width=True,
                        hide_index=True,
                        disabled=["№", "ФИО студента"],
                        column_config=grade_col_config,
                        key="grade_editor",
                    )

                    meta_columns = ["Поле"] + date_cols
                    meta_data = [
                        {
                            "Поле": "Тема занятия",
                            **{
                                pd.to_datetime(d).strftime("%d.%m"): topic_map.get(
                                    d, ""
                                )
                                for d in dates
                            },
                        },
                        {
                            "Поле": "Домашнее задание",
                            **{
                                pd.to_datetime(d).strftime("%d.%m"): hw_map.get(d, "")
                                for d in dates
                            },
                        },
                    ]
                    meta_df = pd.DataFrame(meta_data, columns=meta_columns)

                    st.markdown("**Темы и домашние задания**")
                    edited_meta_df = st.data_editor(
                        meta_df,
                        use_container_width=True,
                        hide_index=True,
                        disabled=["Поле"],
                        column_config={
                            col: st.column_config.TextColumn(col) for col in date_cols
                        },
                        key="meta_editor",
                    )

                    st.divider()
                    excel_buf = generate_excel(
                        disc_name,
                        group_name,
                        teacher_name,
                        dates,
                        students,
                        lesson_map,
                        topic_map,
                        hw_map,
                        conn,
                    )
                    safe_name = f"{disc_name}_{group_name}".replace(" ", "_")
                    st.download_button(
                        label="⬇️ Экспорт в Excel",
                        data=excel_buf,
                        file_name=f"{safe_name}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                    st.divider()

                    with st.expander("Добавить урок"):
                        new_date = st.date_input("Дата урока", date.today())
                        dstr = new_date.strftime("%Y-%m-%d")
                        if dstr in dates:
                            st.warning("Урок на эту дату уже существует")
                        else:
                            topic = st.text_input("Тема занятия")
                            hw = st.text_area("Домашнее задание (необязательно)", "")
                            if st.button("Создать урок"):
                                c.execute(
                                    "INSERT INTO lessons (group_discipline_id, date, homework, topic) VALUES (?,?,?,?)",
                                    (assignment_id, dstr, hw, topic),
                                )
                                conn.commit()
                                st.success("Урок добавлен!")
                                st.rerun()

                    if st.button("Сохранить все изменения"):
                        changed = False
                        for i in range(len(students)):
                            sid = students[i][0]
                            for d in dates:
                                col = pd.to_datetime(d).strftime("%d.%m")
                                old_val = (
                                    str(grade_df.iloc[i][col]).strip()
                                    if pd.notna(grade_df.iloc[i][col])
                                    else ""
                                )
                                new_val = (
                                    str(edited_grade_df.iloc[i][col]).strip()
                                    if pd.notna(edited_grade_df.iloc[i][col])
                                    else ""
                                )
                                if new_val != old_val:
                                    changed = True
                                    if new_val in ["", "—", "None"]:
                                        c.execute(
                                            "DELETE FROM grades WHERE lesson_id=? AND student_id=?",
                                            (lesson_map[d], sid),
                                        )
                                    else:
                                        c.execute(
                                            "INSERT OR REPLACE INTO grades (lesson_id, student_id, grade) VALUES (?,?,?)",
                                            (lesson_map[d], sid, new_val),
                                        )

                        for d in dates:
                            col = pd.to_datetime(d).strftime("%d.%m")
                            old_topic = topic_map.get(d, "")
                            new_topic = (
                                str(edited_meta_df.iloc[0][col]).strip()
                                if pd.notna(edited_meta_df.iloc[0][col])
                                else ""
                            )
                            if new_topic != old_topic:
                                changed = True
                                c.execute(
                                    "UPDATE lessons SET topic=? WHERE id=?",
                                    (new_topic, lesson_map[d]),
                                )

                            old_hw = hw_map.get(d, "")
                            new_hw = (
                                str(edited_meta_df.iloc[1][col]).strip()
                                if pd.notna(edited_meta_df.iloc[1][col])
                                else ""
                            )
                            if new_hw != old_hw:
                                changed = True
                                c.execute(
                                    "UPDATE lessons SET homework=? WHERE id=?",
                                    (new_hw, lesson_map[d]),
                                )

                        if changed:
                            conn.commit()
                            st.success("Все изменения сохранены!")
                            st.rerun()
                        else:
                            st.info("Изменений нет")

    # ======================= АДМИН ПАНЕЛЬ =======================
    if user["role"] == "admin" and page != "Журнал":
        if page == "Пользователи":
            st.subheader("Создать пользователя")
            new_username = st.text_input("Логин")
            new_password = st.text_input("Пароль", type="password")
            new_full_name = st.text_input("ФИО")
            new_role = st.selectbox("Роль", ["student", "teacher", "parent", "admin"])

            new_group_id = None
            selected_children_ids = []

            if new_role == "student":
                groups_all = c.execute("SELECT id, name FROM groups").fetchall()
                if groups_all:
                    gname = st.selectbox("Группа", [g[1] for g in groups_all])
                    new_group_id = next(
                        (g[0] for g in groups_all if g[1] == gname), None
                    )

            elif new_role == "parent":
                students_all = c.execute("""
                    SELECT u.id, u.full_name, COALESCE(g.name, '—') as group_name
                    FROM users u
                    LEFT JOIN groups g ON g.id = u.group_id
                    WHERE u.role = 'student'
                    ORDER BY u.full_name
                """).fetchall()

                if students_all:
                    student_options = [f"{row[1]} ({row[2]})" for row in students_all]
                    selected_display = st.multiselect(
                        "Привязать детей", student_options
                    )

                    for disp in selected_display:
                        for row in students_all:
                            if f"{row[1]} ({row[2]})" == disp:
                                selected_children_ids.append(row[0])
                                break

            if st.button("Создать пользователя"):
                if not new_username or not new_password or not new_full_name:
                    st.error("Заполните все поля")
                else:
                    h = hashlib.sha256(new_password.encode()).hexdigest()
                    try:
                        c.execute(
                            "INSERT INTO users (username,password,full_name,role,group_id) VALUES (?,?,?,?,?)",
                            (new_username, h, new_full_name, new_role, new_group_id),
                        )
                        conn.commit()
                        new_user_id = c.lastrowid

                        if new_role == "parent" and selected_children_ids:
                            for child_id in selected_children_ids:
                                c.execute(
                                    "INSERT OR IGNORE INTO parent_students (parent_id, student_id) VALUES (?,?)",
                                    (new_user_id, child_id),
                                )
                            conn.commit()

                        st.success(f"Пользователь {new_username} успешно создан!")
                        st.rerun()
                    except sqlite3.IntegrityError:
                        st.error("Логин уже занят")

        elif page == "Группы":
            st.subheader("Группы")
            name = st.text_input("Название новой группы")
            if st.button("Добавить группу") and name:
                c.execute("INSERT OR IGNORE INTO groups (name) VALUES (?)", (name,))
                conn.commit()
                st.rerun()
            groups = c.execute("SELECT id, name FROM groups ORDER BY name").fetchall()
            if groups:
                st.dataframe(
                    pd.DataFrame(groups, columns=["ID", "Название"]),
                    use_container_width=True,
                    hide_index=True,
                )

        elif page == "Предметы":
            st.subheader("Предметы")
            name = st.text_input("Название нового предмета")
            if st.button("Добавить предмет") and name:
                c.execute("INSERT INTO disciplines (name) VALUES (?)", (name,))
                conn.commit()
                st.rerun()
            discs = c.execute(
                "SELECT id, name FROM disciplines ORDER BY name"
            ).fetchall()
            if discs:
                st.dataframe(
                    pd.DataFrame(discs, columns=["ID", "Название"]),
                    use_container_width=True,
                    hide_index=True,
                )

        elif page == "Назначение":
            st.subheader("Назначить преподавателя")
            groups = c.execute("SELECT id, name FROM groups ORDER BY name").fetchall()
            discs = c.execute(
                "SELECT id, name FROM disciplines ORDER BY name"
            ).fetchall()
            teachers = c.execute(
                "SELECT id, full_name FROM users WHERE role='teacher' ORDER BY full_name"
            ).fetchall()

            if groups and discs and teachers:
                g = st.selectbox("Группа", [x[1] for x in groups])
                d = st.selectbox("Предмет", [x[1] for x in discs])
                t = st.selectbox("Преподаватель", [x[1] for x in teachers])

                gid = next(x[0] for x in groups if x[1] == g)
                did = next(x[0] for x in discs if x[1] == d)
                tid = next(x[0] for x in teachers if x[1] == t)

                if st.button("Назначить"):
                    c.execute(
                        "INSERT OR IGNORE INTO group_disciplines (group_id, discipline_id, teacher_id) VALUES (?,?,?)",
                        (gid, did, tid),
                    )
                    conn.commit()
                    st.success("Назначено успешно!")
            else:
                st.warning("Сначала создайте группы, предметы и преподавателей")

conn.close()
